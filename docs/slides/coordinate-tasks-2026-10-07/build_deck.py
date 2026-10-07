"""Build an editable technical deck from pinned, independently checked evidence.

Run with Python plus python-pptx and numpy installed:
    python docs/slides/coordinate-tasks-2026-10-07/build_deck.py

All charts, diagrams and labels are editable PowerPoint objects. The only
geometry visual is projected from a real task's displayed PDB coordinates.
"""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "evidence.json").read_text())
M, C, T = DATA["manifest"], DATA["context"], DATA["teacher"]
URL = DATA["source_urls"]
OUT = ROOT / "pdbthink-coordinate-tasks-2026-10-07.pptx"

INK = "142E35"
NIGHT = "10282F"
WHITE = "FFFFFF"
PAPER = "F5F4EF"
TEAL = "067D79"
MINT = "91E1CD"
PALE = "DFEDE7"
ORANGE = "D97941"
AMBER = "F3BD74"
MUTED = "586E72"
SOFT = "A8BBC0"
LINE = "D4DDDA"
CORAL = "BE5846"
SANS = "Arial"
MONO = "Liberation Mono"

NAMES = {
    "P01": "Chain identifiers", "P02": "Residue counting", "P03": "Atom coordinates",
    "G01": "Atom distance", "G02": "Nearest eligible atom", "G03": "Nearest residue",
    "G04": "Steric clash", "S01": "Salt bridges", "S02": "Phosphorylation",
    "S03": "Solvent exposure", "S04": "Secondary structure", "S05": "Fold class",
    "S06": "Ligand contacts", "S07": "Metal coordination", "S08": "Disulfide partner",
    "S09": "Side-chain rotamer", "I01": "Interface contacts", "N01": "Shared contact",
    "T01": "Two-state contacts",
}
ORDER = list(NAMES)
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
prs.core_properties.title = "PDBThink Coordinate Tasks"
prs.core_properties.subject = "Technical overview for Snowball evaluation and post-training"
prs.core_properties.author = "Open Athena"
prs.core_properties.keywords = "PDBThink, Snowball, coordinate reasoning, RLVR, SFT"
prs.core_properties.comments = "Pinned task release v1.3.0; historical teacher release v1.0.0 on tasks v1.2.0."
SLIDE_META = []


def rgb(value):
    return RGBColor.from_string(value)


def box(s, x, y, w, h, fill=WHITE, stroke=None, radius=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if stroke:
        sh.line.color.rgb = rgb(stroke)
        sh.line.width = Pt(0.7)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    if radius:
        sh.adjustments[0] = 0.035
    return sh


def txt(s, value, x, y, w, h, size=18, color=INK, bold=False, font=SANS,
        align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, link=None):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, line in enumerate(str(value).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
    if link:
        sh.click_action.hyperlink.address = link
    return sh


def line(s, x1, y1, x2, y2, color=LINE, width=1.0):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color)
    sh.line.width = Pt(width)
    return sh


def dot(s, x, y, radius=0.05, fill=TEAL, stroke=None):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x-radius), Inches(y-radius),
                            Inches(radius*2), Inches(radius*2))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if stroke:
        sh.line.color.rgb = rgb(stroke)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def pill(s, label, x, y, w, fill=PALE, color=TEAL):
    box(s, x, y, w, 0.32, fill, radius=True)
    txt(s, label, x+0.09, y+0.06, w-0.18, 0.22, 10, color, True)


def base(title, section, subtitle="", *, dark=False, source="tasks", notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NIGHT if dark else PAPER)
    fg, sub = (WHITE, SOFT) if dark else (INK, MUTED)
    txt(s, section.upper(), 0.58, 0.32, 11.9, 0.24, 10, MINT if dark else TEAL, True)
    txt(s, title, 0.58, 0.79, 12.12, 0.72, 31, fg, True)
    if subtitle:
        txt(s, subtitle, 0.6, 1.56, 12.05, 0.48, 14, sub)
    n = len(prs.slides)
    line(s, 0.58, 7.08, 12.72, 7.08, "365057" if dark else LINE, 0.6)
    txt(s, "OPEN ATHENA  /  PDBTHINK COORDINATE TASKS", 0.6, 7.2, 5.6, 0.18, 8.4, sub)
    label = {"tasks": "Task release v1.3.0", "manifest": "Task manifest v1.3.0", "context": "Snowball native-context audit",
             "teacher": "Historical GLM traces · tasks v1.2.0", "teacher_summary": "GLM results · tasks v1.2.0",
             "corrections": "v1.3.0 correction audit", "code": "Frozen generator source"}.get(source, source)
    txt(s, label, 6.9, 7.18, 5.05, 0.23, 9, sub, align=PP_ALIGN.RIGHT, link=URL.get(source))
    txt(s, f"{n:02d}", 12.2, 7.18, 0.5, 0.22, 9, sub, align=PP_ALIGN.RIGHT)
    full_notes = notes + "\n\nEvidence: " + URL.get(source, URL["tasks"])
    s.notes_slide.notes_text_frame.text = full_notes
    SLIDE_META.append({"number":n, "title":title, "section":section, "notes":full_notes})
    return s


def callout(s, heading, body, x, y, w, h, *, dark=False, accent=TEAL):
    box(s, x, y, w, h, "1B3941" if dark else WHITE, radius=True)
    box(s, x, y, 0.045, h, accent)
    txt(s, heading, x+0.21, y+0.18, w-0.4, 0.45, 19, MINT if dark else accent, True)
    txt(s, body, x+0.21, y+0.76, w-0.43, h-0.88, 16, WHITE if dark else INK)


def molecule(s, x, y, w, h, *, dark=False, labels=False):
    atoms = DATA["example_atoms"]
    xyz = np.array([a["xyz"] for a in atoms])
    centred = xyz-xyz.mean(axis=0)
    _, _, vt = np.linalg.svd(centred, full_matrices=False)
    uv = centred @ vt.T
    uv = np.stack([uv[:,1]+0.23*uv[:,2], -uv[:,0]+0.15*uv[:,2]], axis=1)
    uv -= uv.min(axis=0)
    scale = min((w-0.8)/max(uv[:,0].max(),1), (h-0.7)/max(uv[:,1].max(),1))
    uv *= scale
    uv[:,0] += x+(w-uv[:,0].max())/2
    uv[:,1] += y+(h-uv[:,1].max())/2
    soft = "3D646D" if dark else "BECFCA"
    for i in range(len(atoms)):
        for j in range(i):
            if np.linalg.norm(xyz[i]-xyz[j]) < 1.85:
                line(s, *uv[i], *uv[j], soft, 1.1)
    targets = [i for i,a in enumerate(atoms) if a["name"] == "CA" and a["number"] in [44,48]]
    line(s, *uv[targets[0]], *uv[targets[1]], ORANGE if not dark else AMBER, 2.4)
    for i in np.argsort(centred[:,2]):
        highlight = i in targets
        color = (AMBER if dark else ORANGE) if highlight else (MINT if dark else TEAL)
        dot(s, *uv[i], radius=0.075 if highlight else 0.031, fill=color)
    if labels:
        for i in targets:
            a=atoms[i]
            txt(s, f"P:G{a['number']}:CA", uv[i,0]+0.14, uv[i,1]-0.09, 1.48, 0.3, 14,
                WHITE if dark else INK, True)
    return uv


# 01 — Title.
s = base("", "Open Athena · technical briefing · 7 October 2026", dark=True,
         notes="This deck describes the coordinate-task training and evaluation release, not the separate sequence-prediction tracks. All counts are pinned. No Snowball learning improvement is claimed: the final slides propose experiments.")
SLIDE_META[-1]["title"] = "PDBThink Coordinate Tasks"
txt(s, "PDBThink", 0.68, 1.32, 7.7, 0.95, 54, WHITE, True)
txt(s, "Coordinate Tasks", 0.68, 2.29, 8.15, 0.85, 42, MINT, True)
txt(s, "A verified task pool for Snowball\nevaluation, RL and teacher-generated SFT", 0.72, 3.42, 7.1, 1.06, 23, WHITE)
pill(s, "TASK RELEASE  v1.3.0", 0.72, 4.8, 2.0, "244750", MINT)
for x, val, label in [(0.72,"100,000","coordinate tasks"),(3.35,"19","problem families"),(5.98,"28,045","8K-headroom train tasks")]:
    txt(s, val, x, 5.55, 2.55, 0.58, 33, WHITE, True)
    txt(s, label, x, 6.21, 2.7, 0.39, 12.5, SOFT)
molecule(s, 8.55, 1.42, 3.67, 4.95, dark=True)
txt(s, "Geometry from an actual released task", 8.35, 6.57, 4.05, 0.25, 10, SOFT, align=PP_ALIGN.CENTER)

# 02 — Scientific contract and uses.
s = base("The answer must follow from the displayed coordinates", "01 / Scientific contract",
         "One structure (or two states), one question, one tool-free response, one deterministic reward.",
         notes="The model sees sanitised coordinate text and a question. Gold is computed from those exact coordinates after rotation, rounding and any permitted crop. The scorer runs outside the model. RLVR means reinforcement learning with verifiable rewards. These tasks do not ask for structure prediction from sequence.")
steps=[("Coordinate text", "Sanitised PDB records"),("Model", "Tools explicitly disabled"),("Final answer", "Typed answer schema"),("Verifier", "Exact 0/1 reward")]
for i,(title,body) in enumerate(steps):
    x=0.65+i*3.15
    box(s,x,2.35,2.8,1.6,WHITE,radius=True)
    txt(s,f"0{i+1}",x+0.16,2.52,0.5,0.35,12,TEAL,True)
    txt(s,title,x+0.16,2.97,2.5,0.36,21,INK,True)
    txt(s,body,x+0.16,3.51,2.5,0.29,12.5,MUTED)
    if i<3:txt(s,"→",x+2.87,2.94,0.26,0.5,23,TEAL,True)
for x,title,body in [(0.65,"Evaluation","Measure held-out coordinate reasoning by family."),(4.85,"RL","Use executable rewards without a model judge."),(9.05,"SFT","Keep verified teacher answers that fit the student.")]:
    callout(s,title,body,x,4.67,3.65,1.65)

# 03 — Taxonomy.
s = base("Nineteen families span parsing, geometry and structure", "02 / Task coverage",
         "All active coordinate families are present. Sequence prediction and retired MECH tasks are excluded.",
         notes="The groups here are for presentation. S05 is a global fold-class task; S02 identifies a chemical modification. The exact task definitions live in the pinned source snapshot. We reuse those definitions but generate new semantic tasks on a benchmark-disjoint source pool.")
groups=[("PARSING",["P01","P02","P03"]),("GEOMETRY",["G01","G02","G03","G04"]),("STRUCTURE",["S02","S03","S04","S05","S09"]),("INTERACTIONS",["S01","S06","S07","S08"]),("RELATIONSHIPS",["I01","N01","T01"])]
for i,(label,codes) in enumerate(groups):
    x=0.62+i*2.55
    box(s,x,2.25,2.35,4.34,WHITE,radius=True)
    box(s,x,2.25,2.35,0.055,TEAL if i%2==0 else ORANGE)
    txt(s,label,x+0.16,2.52,2.05,0.3,12,TEAL,True)
    for j,code in enumerate(codes):
        yy=3.09+j*0.65
        txt(s,code,x+0.15,yy,0.51,0.25,12,TEAL,True,font=MONO)
        txt(s,NAMES[code],x+0.69,yy-0.01,1.5,0.52,14,INK)

# 04 — Verified worked example.
e=DATA["example"]
a=next(a for a in DATA["example_atoms"] if a["name"]=="CA" and a["number"]==44)
b=next(a for a in DATA["example_atoms"] if a["name"]=="CA" and a["number"]==48)
s = base("Read two atoms; calculate their distance", "03 / A real released task",
         "G01 · 54 displayed atoms · only the two relevant PDB records are excerpted below.",
         notes=f"Task: {e['task_id']}. Train split, task release v1.3.0. Full original prompt, hidden gold and checksums are in evidence.json. The example was selected from the first train shard for a compact visual, not as a representative random sample. Independently parsed fixed-width columns and math.dist gave {DATA['example_distance']:.12f} Å; packaged gold is 10.903 Å. Answer 10.90 satisfies tolerance ±0.02 Å. The complete model input has 54 atoms; this slide is explicitly an excerpt, not a new evaluation prompt.")
box(s,0.65,2.2,8.2,1.65,NIGHT,radius=True)
txt(s,"EXCERPT FROM THE MODEL INPUT",0.84,2.37,7.8,0.23,10,MINT,True)
txt(s,a["line"][:54]+"\n"+b["line"][:54],0.84,2.86,7.82,0.71,14,WHITE,font=MONO)
txt(s,"P:G44:CA",0.78,4.13,1.85,0.35,20,TEAL,True)
txt(s,"(−0.514,  7.195,  2.794)",2.64,4.16,5.3,0.32,18,INK,font=MONO)
txt(s,"P:G48:CA",0.78,4.76,1.85,0.35,20,TEAL,True)
txt(s,"(−0.079, −3.046, −0.922)",2.64,4.79,5.3,0.32,18,INK,font=MONO)
txt(s,"d = √(0.435² + 10.241² + 3.716²) = 10.903 Å",0.78,5.57,8.0,0.5,22,INK,True)
pill(s,"FINAL: 10.90     →     reward = 1",0.78,6.34,4.02,PALE,TEAL)
molecule(s,9.02,2.27,3.55,3.92,labels=False)
txt(s,"The orange line joins the queried atoms.",9.0,6.35,3.57,0.37,11.5,MUTED)

# 05 — Generation and validation.
s = base("Check gold against the displayed coordinates", "04 / Generation pipeline",
         "The benchmark’s family generators and operational definitions are reused on new source material.", source="code",
         notes="Only acquisition touches the network. Stable identifiers determine seeds. Question proposals identify parameters; they are not accepted as gold. The oracle runs on the final rotated, rounded and permitted-cropped structures. A separate text parser reconstructs the displayed PDB and invokes the oracle again. Ambiguities and failed margins are rejected into the ledger. A rotated copy alone never creates a new semantic instance. All 100,000 packaged tasks passed coordinate recomputation and solution verification in the release audit.")
stages=[("1", "Acquire & exclude", "Frozen PDB pool; audit every protein partner."),("2", "Propose & render", "Choose a question; sanitise, transform and round."),("3", "Recompute gold", "Run the family oracle on the displayed geometry."),("4", "Parse & check", "Read the emitted PDB text and rerun the oracle.")]
for i,(num,title,body) in enumerate(stages):
    x=0.65+i*3.14
    dot(s,x+0.26,2.53,0.22,TEAL)
    txt(s,num,x+0.08,2.39,0.36,0.28,17,WHITE,True,align=PP_ALIGN.CENTER)
    if i<3:line(s,x+0.53,2.53,x+2.98,2.53,LINE,2)
    txt(s,title,x,3.08,2.92,0.7,21,INK,True)
    txt(s,body,x,3.99,2.7,1.22,18,MUTED)
box(s,0.65,5.77,12.0,0.84,PALE,radius=True)
txt(s,"100,000 / 100,000",0.86,5.93,3.1,0.48,27,TEAL,True)
txt(s,"packaged coordinate checks and oracle-solution checks passed",4.1,6.03,8.2,0.35,17,INK)

# 06 — Separation and split discipline.
s = base("Hold out related proteins, not random task rows", "05 / Benchmark separation",
         "The release excludes the frozen PDBThink source inventory and keeps related sources in one split.", source="manifest",
         notes="Exclusions cover entry IDs, exact full-entity and observed-chain protein sequences, and RCSB 30% sequence clusters associated with configured benchmark sources. Every protein partner is audited. Entries connected by sequences, clusters or two-state pairs form one source group. Groups are deterministically assigned to a target 90/5/5 split, giving the displayed actual counts. All families occur in all three splits. Zero audited overlap is relative to that frozen inventory; it is not a claim of freedom from foundation-model pretraining.")
for yy,label,detail in [(2.28,"Entry identity","Every configured benchmark source"),(3.27,"Protein sequence","Full-entity and observed-chain sequences"),(4.26,"Related-source clusters","RCSB 30% sequence clusters")]:
    box(s,0.65,yy,5.76,0.83,WHITE,radius=True)
    txt(s,"0",0.84,yy+0.1,0.53,0.52,31,TEAL,True)
    txt(s,label,1.51,yy+0.09,4.65,0.29,17,INK,True)
    txt(s,detail,1.51,yy+0.47,4.65,0.23,12,MUTED)
txt(s,"Audited intersections with the benchmark",0.7,5.38,5.56,0.45,14,MUTED)
txt(s,"1,867 source groups",7.03,2.26,5.32,0.54,28,INK,True)
split_items=[("Train",91154,TEAL),("Validation",4411,ORANGE),("Test",4435,INK)]
for i,(label,n,col) in enumerate(split_items):
    y=3.15+i*0.77
    dot(s,7.18,y+0.16,0.07,col)
    txt(s,label,7.4,y,2.35,0.4,20,INK)
    txt(s,f"{n:,}",10.15,y,2.15,0.4,24,col,True,align=PP_ALIGN.RIGHT)
box(s,0.65,6.11,12.0,0.57,PALE,radius=True)
txt(s,"This establishes benchmark separation; it does not establish absence from model pretraining.",0.86,6.25,11.61,0.28,15,INK)

# 07 — Family imbalance.
s = base("The 100,000 tasks are deliberately broad, but uneven", "06 / Dataset composition",
         "Nine abundant families account for 88.2% of tasks. Rare admissible cases limit the remaining families.", source="manifest",
         notes="Counts are semantic tasks, not independent proteins. The pool has 2,671 PDB entries in 1,867 source groups. The expanded release increases questions per source, not tenfold independent structures. The nine large families sum to 88,152 tasks. Use family-aware sampling and source-group-aware evaluation; a row-weighted average can obscure low-frequency families.")
sorted_codes=sorted(ORDER,key=lambda k:-M['family_counts'][k])
for i,code in enumerate(sorted_codes):
    yy=2.12+i*0.23
    txt(s,code,0.73,yy,0.58,0.19,10.7,INK,True,font=MONO)
    box(s,1.51,yy+0.025,5.45*M['family_counts'][code]/10000,0.145,TEAL if i<9 else ORANGE)
    txt(s,f"{M['family_counts'][code]:,}",7.05,yy-0.005,0.91,0.2,10.7,INK,align=PP_ALIGN.RIGHT)
for x in [0,5000,10000]:
    txt(s,f"{x:,}",1.51+5.45*x/10000-0.22,6.62,0.65,0.23,10,MUTED,align=PP_ALIGN.CENTER)
line(s,1.51,6.52,6.96,6.52,LINE)
callout(s,"2,671 PDB entries","Multiple questions may use the same source structure.",8.54,2.28,4.08,1.61)
callout(s,"232 steric-clash tasks","Compare with 9,795 tasks in each of several abundant families.",8.54,4.11,4.08,1.82,accent=ORANGE)
txt(s,"Report family-level results alongside aggregate accuracy.",8.56,6.16,4.05,0.62,17,INK,True)

# 08 — Harbor packaging and isolation.
s = base("Harbor packages a tool-free model task", "07 / Execution contract",
         "Task Trove-style Parquet rows contain task_binary and a separate solution_binary.",
         notes="task_binary contains the instruction, prompt.json, task.toml, Dockerfile and hidden test verifier. Gold and tests are evaluator-only; they must not be forwarded to the model. solution_binary contains the oracle final answer and is used only for offline mechanical checks. CoordinateNoToolsAgent sends system/user messages with tools explicitly disabled, retains raw responses, audits tool events and writes answer.txt on the host. Generic terminal agents violate this protocol. Contact-set reward requires the full correct set; partial-set F1 is diagnostic only.")
box(s,0.65,2.23,5.25,3.9,NIGHT,radius=True)
txt(s,"MODEL INPUT",0.9,2.5,4.75,0.27,11,MINT,True)
txt(s,"prompt.json",0.9,3.1,4.65,0.42,27,WHITE,True,font=MONO)
txt(s,"System instructions\nExact coordinate question\nNo gold, provenance or source IDs",0.9,3.84,4.67,1.47,19,WHITE)
txt(s,"One response → FINAL field",0.9,5.47,4.65,0.33,17,MINT,True)
txt(s,"→",6.08,3.83,0.68,0.65,37,TEAL,True)
box(s,6.93,2.23,5.71,3.9,WHITE,radius=True)
txt(s,"HOST / EVALUATOR",7.2,2.5,5.05,0.27,11,TEAL,True)
txt(s,"Bundled deterministic verifier",7.2,3.1,5.05,0.7,24,INK,True)
txt(s,"Typed parsing + numerical tolerances\nExact correctness → reward 0 or 1\nPartial-credit diagnostics kept separately",7.2,4.0,5.08,1.41,18,INK)
pill(s,"ORACLE SOLUTION REMAINS SEPARATE",7.2,5.6,3.72)
txt(s,"For evaluation, allocate the largest supported output budget that fits the exact native-token prompt.",0.78,6.47,11.75,0.39,16,MUTED)

# 09 — v1.3 fixes.
s = base("v1.3 makes the scoring boundary and clash rule explicit", "08 / Current release",
         "All 100,000 task identities, displayed coordinates, gold answers and grouped splits are preserved.", source="corrections",
         notes="Numeric scoring 1.1.0 compares exact rational differences of the decimal representations of parsed values, retaining the original tolerance. The numeric parser itself still parses floats. The teacher audit found 11 affected responses across two tasks. Prompt v5 states G04's overlap/radii, covalent-neighbour and metal exclusions, unconditional sulfur SG–SG exclusion, and residue-pair ranking. All 232 G04 prompts changed. The correction audit tested the actual packaged verifier in Docker. Historical GLM v1.0.0 traces remain on task v1.2.0 and are not evaluations of the revised G04 question.")
callout(s,"Inclusive numeric tolerance","−1.710 vs −1.711 should pass ±0.001.\n\nScorer 1.1.0 removes binary-subtraction boundary errors; tolerance is unchanged.",0.65,2.3,5.81,3.29)
callout(s,"Steric-clash exclusions","The question now explicitly excludes every SG–SG pair.\n\nPrompt v5 clarifies all 232 G04 tasks without changing their gold answers.",6.72,2.3,5.91,3.29,accent=ORANGE)
box(s,0.65,5.94,11.98,0.82,INK,radius=True)
txt(s,"Version boundary",0.88,6.13,2.35,0.36,18,MINT,True)
txt(s,"Published GLM traces retain their v1.2 prompts and historical rewards.",3.35,6.15,9.03,0.4,17,WHITE)

# 10 — Context budget.
tr=C['splits']['train']
s = base("28,045 training tasks leave 8K for reasoning and output", "09 / Snowball context", dark=True,
         subtitle="Exact token counts use Snowball’s pinned tokenizer and thinking-enabled chat template.",source="context",
         notes="Snowball checkpoint: open-athena/Snowball-67B-A2B-5.7T-Mixed-RLVR-Step38 at cfc1d845dae89b067cdc7250d0164abefa5a69cf. The native window is 32,768 tokens. 8K headroom means prompt plus generation prefix is at most 24,576 tokens. This is a cohort filter, not an imposed generation cap. Actual SFT eligibility includes the teacher completion, template delimiters and end token. Across all splits, 30,418 tasks qualify: 28,045 train, 1,003 validation and 1,370 test. T01 has zero 8K-headroom tasks. The train cohort covers 18 families; all-split eligible tasks span 487 source groups.")
box(s,0.74,2.31,8.92,0.74,"22434C",radius=True)
box(s,0.74,2.31,6.69,0.74,TEAL)
box(s,7.43,2.31,2.23,0.74,AMBER)
txt(s,"Prompt ≤24,576 tokens",0.99,2.54,5.8,0.29,18,WHITE,True)
txt(s,"8,192 reserve",7.62,2.56,1.9,0.28,15,NIGHT,True)
txt(s,"32,768 total",10.03,2.52,2.3,0.36,20,WHITE,True)
values=[("All training tasks",tr['total']), ("Prompt alone fits",tr['prompt_under_32k']), ("8K headroom",tr['room_for_8192']), ("16K headroom",tr['room_for_16384'])]
for i,(label,n) in enumerate(values):
    y=3.63+i*0.66
    txt(s,label,0.77,y,3.14,0.3,17,WHITE)
    box(s,4.07,y+0.01,5.63*n/tr['total'],0.27,MINT if i==2 else "4D7881")
    txt(s,f"{n:,}",10.03,y-0.05,2.14,0.41,23,MINT if i==2 else WHITE,True)
txt(s,"Held-out 8K cohort: 1,003 validation + 1,370 test",0.78,6.48,8.28,0.36,17,SOFT)
pill(s,"18 FAMILIES · NO T01",9.73,6.42,2.59,"244750",MINT)

# 11 — Teacher setup.
s = base("GLM samples independently until success or ten attempts", "10 / Teacher traces",
         "Historical run: GLM-5.3 on the 28,045 eligible training tasks from task release v1.2.0.",source="teacher",
         notes="The task cohort is training-only and spans 438 source groups. Tools are explicitly disabled; the endpoint uses tools:null and tool_choice:none. High reasoning effort, temperature 1.0, top-p 0.95. Each retry uses the identical prompt and a fresh deterministic seed, with no previous answers, gold or verifier feedback. Stop at first native-verifier success or ten scored attempts. Transport errors are separate. Teacher generation receives its full remaining served context of 262,144 tokens, rather than an 8,192-token cap. The teacher's immutable served weight revision is unavailable; the tokenizer, task release and request policy are pinned.")
for i,(title,body) in enumerate([("Sample", "Frozen prompt\nFresh seed\nNo tools"),("Score", "Task’s bundled verifier\nCorrectness recorded\nRaw response retained"),("Select or retry", "Stop at first success\nOtherwise retry unchanged\nMaximum ten attempts")]):
    x=0.65+i*4.15
    callout(s,title,body,x,2.43,3.75,2.87)
    if i<2:txt(s,"→",x+3.82,3.59,0.28,0.43,25,TEAL,True)
box(s,0.65,5.73,12.05,0.92,PALE,radius=True)
txt(s,"8K selects the student cohort; it does not cap the teacher.",0.91,5.92,11.45,0.45,23,TEAL,True)

# 12 — Adaptive retry result.
success=Counter()
for family in T['families'].values():success.update(family['first_correct'])
cumulative=np.cumsum([success[str(i)] for i in range(1,11)])
s = base("Retries recover 1,993 of the 2,156 first-attempt failures", "11 / Teacher yield",
         "The chart tracks recovery among the 2,156 initially failed tasks · 33,003 total scored attempts.",source="teacher_summary",
         notes="The curve denominator is the 2,156 tasks that failed the first attempt; the right-hand headline percentages instead use all 28,045 tasks in the eligible training cohort. Success counts by first successful attempt: "+json.dumps(dict(success))+". These are observed adaptive-run yields, not a fixed-sample pass@k estimator. First-try accuracy is 92.31%, by three attempts 98.27%, by ten 99.42%. By attempts two, three and ten, retries recovered 1,388, 1,670 and 1,993 initially failed tasks respectively. 163 tasks exhausted ten attempts under the frozen verifier. Eleven numeric boundary rejections across two tasks and an underspecified clash rule were later audited; historical scores remain unchanged. Correct final answers do not verify every reasoning step, and retries can find categorical labels by chance.")
x0,y0,w,h=0.95,2.54,7.56,3.57
for pct in [0,25,50,75,100]:
    yy=y0+h-h*pct/100
    line(s,x0,yy,x0+w,yy,LINE,.6)
    txt(s,f"{pct}%",0.32,yy-0.1,.5,.21,10,MUTED,align=PP_ALIGN.RIGHT)
pts=[]
recovery=cumulative[1:]-T['first_try']
initial_failures=T['task_count']-T['first_try']
for i,n in enumerate(recovery):
    xx=x0+w*i/8; yy=y0+h-h*n/initial_failures
    pts.append((xx,yy))
    txt(s,str(i+2),xx-.17,y0+h+.16,.34,.2,11,MUTED,align=PP_ALIGN.CENTER)
for a,b in zip(pts,pts[1:]):line(s,*a,*b,TEAL,2.8)
for xx,yy in pts:dot(s,xx,yy,.052,TEAL)
for i,label,dy in [(0,"64.4%",.24),(1,"77.5%",.26),(8,"92.4%",.22)]:
    xx,yy=pts[i]; txt(s,label,xx-(.95 if i==8 else .18),yy+dy,1.03,.29,15,TEAL,True,align=PP_ALIGN.CENTER)
txt(s,"Total attempts per task",3.44,6.5,3.71,.26,13,MUTED,align=PP_ALIGN.CENTER)
callout(s,"92.31% → 99.42%","Overall success: first try → within ten attempts.",9.0,2.35,3.65,1.79)
callout(s,"163 remain unsolved","Ten failed attempts are not proof that a task is impossible.",9.0,4.39,3.65,1.91,accent=ORANGE)

# 13 — Harder family profile.
s = base("Lower first-try success is concentrated in a few families", "12 / Family-level results",
         "Selected lower-first-try families. Small cohorts are shown explicitly; full counts are in the appendix.",source="teacher_summary",
         notes="Families are selected for lower first-attempt success, not as a representative subset. All results are historical task-v1.2.0 native-verifier results. Interface contacts have only 12 tasks, ligand contacts 53, clashes 58 and metal coordination 92; avoid broad conclusions from those denominators. Solvent exposure and secondary structure account for 128 of the 163 exhausted tasks. G04 results predate the clarified SG–SG exclusion. The full family table is included later.")
picked=["I01","S06","S07","G04","S04","S05","S03"]
txt(s,"Family",0.71,2.17,3.9,.28,12,MUTED,True)
txt(s,"Tasks",4.31,2.17,.7,.28,12,MUTED,True,align=PP_ALIGN.RIGHT)
txt(s,"First try → by ten",5.49,2.17,5.44,.28,12,MUTED,True)
txt(s,"Unsolved",11.58,2.17,1.0,.28,12,MUTED,True,align=PP_ALIGN.RIGHT)
for i,code in enumerate(picked):
    f=T['families'][code]; yy=2.75+i*.45
    p1,p10=f['first_try']/f['total'],f['solved']/f['total']
    txt(s,f"{code}  {NAMES[code]}",.73,yy,3.58,.31,15,INK)
    txt(s,f"{f['total']:,}",4.23,yy,.77,.28,14,INK,align=PP_ALIGN.RIGHT)
    line(s,5.49,yy+.13,10.66,yy+.13,LINE,2.5)
    xx1,xx2=5.49+5.17*p1,5.49+5.17*p10
    line(s,xx1,yy+.13,xx2,yy+.13,TEAL,3)
    dot(s,xx1,yy+.13,.064,ORANGE);dot(s,xx2,yy+.13,.064,TEAL)
    txt(s,f"{100*p1:.1f}",xx1-.58,yy-.05,.46,.29,10,CORAL,align=PP_ALIGN.RIGHT)
    txt(s,f"{100*p10:.1f}",xx2+.13,yy-.01,.61,.28,10,TEAL)
    txt(s,str(f['unsolved_after_10']),11.75,yy,.72,.28,15,INK,True,align=PP_ALIGN.RIGHT)
box(s,.65,6.21,11.99,.54,PALE,radius=True)
txt(s,"Solvent exposure + secondary structure account for 128 of the 163 exhausted tasks.",.87,6.35,11.53,.29,16,INK)

# 14 — Student-fit yield.
s = base("A correct teacher answer may still be too long for Snowball", "13 / SFT inventory",
         "The published SFT configuration keeps the first correct response only when the entire conversation fits.",source="teacher_summary",
         notes="All values describe the historical teacher release. Of 28,045 cohort tasks, 27,882 have a native-verifier success. 3,360 first-correct responses exceed Snowball's 32,768-token conversation limit, leaving 24,522 SFT examples. Of those, 22,094 have assistant completion length at most 8,192 Snowball tokens. No trace is truncated. The 8K subset is stricter than whole-conversation fit and is not the teacher's generation cap. 22,862 SFT examples have a separately returned reasoning field; 1,660 are answer-only targets with no invented reasoning. I01 contributes zero SFT examples, G04 twelve, and T01 has no eligible prompts.")
stages=[("Eligible train tasks",28045),("Verifier-correct",27882),("Whole conversation fits",24522),("Assistant ≤8K",22094)]
for i,(label,n) in enumerate(stages):
    yy=2.35+i*.83
    txt(s,label,.74,yy,3.75,.46,19,INK, i>1)
    box(s,4.83,yy+.03,4.59*n/28045,.35,TEAL if i>1 else "A2BCB5")
    txt(s,f"{n:,}",9.83,yy-.03,2.23,.51,29,TEAL if i>1 else INK,True,align=PP_ALIGN.RIGHT)
box(s,.65,6.1,12.0,.65,NIGHT,radius=True)
txt(s,"3,360 correct traces exceed context",.89,6.27,6.2,.32,18,MINT,True)
txt(s,"I01: 0 SFT examples    ·    G04: 12",7.45,6.29,4.88,.31,16,WHITE)

# 15 — Recommended training design.
s = base("Use verified SFT first, then test whether RL adds transfer", "14 / Proposed Snowball experiment",dark=True,
         subtitle="This is a proposed experiment sequence; no Snowball training gain has been measured here.",source="teacher",
         notes="Recommendations, not observed results. First make a separately versioned teacher revision: replay numeric scoring with task v1.3, regenerate G04 against the clarified prompt, and retain the historical release. For an initial SFT run, start from the completion_within_8k subset after that reconciliation, apply exact full-sequence token checks, and use assistant-only loss with the pinned thinking-enabled Snowball template. Compare a row-proportional mixture with a family-aware mixture while retaining parsing anchors. Then compare SFT+RL with SFT-only under matched evaluation and a recorded training budget. Do not select RL families solely from teacher failures; use held-out Snowball validation results. Family and source sampling need attention because tasks share structures.")
cards=[("01", "Reconcile versions", "Rescore numeric cases and regenerate G04 on v1.3; publish a new teacher revision."),
       ("02", "Run a focused SFT pilot", "Use verified, student-fitting traces. Compare natural and family-aware mixtures."),
       ("03", "Add RL on residual errors", "Target Snowball’s measured weaknesses using deterministic rewards and grouped splits.")]
for i,(num,title,body) in enumerate(cards):
    x=.67+i*4.18
    box(s,x,2.41,3.78,3.9,"1B3941",radius=True)
    txt(s,num,x+.2,2.62,1,.43,22,MINT,True)
    txt(s,title,x+.2,3.33,3.36,.88,25,WHITE,True)
    txt(s,body,x+.2,4.52,3.31,1.35,18,WHITE)
txt(s,"Keep the original benchmark and the source-group held-out splits out of training.",.83,6.58,11.8,.3,16,SOFT)

# 16 — Evaluation criteria.
s = base("Judge progress by held-out transfer and complete answers", "15 / Evaluation plan",
         "Compare baseline Snowball, SFT-only and SFT + RL on the same versioned evaluation cohorts.",source="context",
         notes="Recommendations only. Start with the fixed v1.3 8K-headroom cohort: 1,003 validation and 1,370 test tasks. Use validation for mixture and checkpoint choices and reserve test for final comparisons. There are no S06, I01 or T01 tasks in either held-out 8K subset. Across the combined held-out subsets, G04 has only five tasks and S07 only three. Report that coverage rather than an all-19-family headline; acquire new independently grouped short instances to improve it. Exact split/family counts were derived from the release's snowball_context.parquet and saved in evidence.json. Report both macro family accuracy and task-weighted accuracy, group-aware uncertainty, schema/format failures, truncation and output tokens. Evaluations must use the largest supported output budget fitting each exact prompt and keep tool use explicitly disabled. Original PDBThink transfer evaluation is separate from the new held-out task pool. Selection effects in teacher retries and repeated sources prevent interpreting the teacher's 99.42% adaptive yield as the student's target single-attempt accuracy.")
rows=[("Capability", "Exact reward by family; macro and task-weighted accuracy."),
      ("Transfer", "Fixed held-out source groups, plus a separate benchmark readout."),
      ("Efficiency", "Tokens, truncation, format errors and tool-policy compliance."),
      ("Coverage", "No held-out 8K tasks for ligand contacts, interfaces or two-state changes.")]
for i,(a,b) in enumerate(rows):
    yy=2.24+i*.83
    box(s,.66,yy,11.99,.67,WHITE,radius=True)
    txt(s,a,.88,yy+.16,2.42,.33,19,TEAL,True)
    txt(s,b,3.39,yy+.18,8.99,.34,17,INK)
box(s,.66,5.96,11.99,.82,PALE,radius=True)
txt(s,"Decision to make",.9,6.18,2.9,.35,19,TEAL,True)
txt(s,"Does RL improve held-out accuracy beyond SFT at an acceptable token cost?",3.82,6.12,8.53,.57,17,INK)

# 17 — Family appendix with consistent denominators.
s = base("Family inventory and teacher outcomes", "Appendix A / Full counts",
         "Task counts use v1.3.0. Teacher outcomes retain task-v1.2.0 scoring; the 8K training task identities are unchanged.",source="teacher_summary",
         notes="Task-pool totals are all splits in task v1.3.0. Eligible train is the v1.2.0 teacher cohort, whose task identities are identical to the v1.3.0 8K training cohort. First and by-ten percentages are historical native-verifier outcomes. SFT counts are first-correct conversations fitting Snowball. The teacher source summary and task manifest supply every cell. Avoid conflating these denominators.")
xs=[.72,4.7,6.24,7.96,9.42,11.05]
ws=[3.78,1.22,1.35,1.13,1.2,1.35]
headers=["Family", "All tasks", "8K train", "First try", "By ten", "SFT fits"]
for x,w,h in zip(xs,ws,headers):txt(s,h,x,2.1,w,.24,11,MUTED,True,align=PP_ALIGN.LEFT if x==xs[0] else PP_ALIGN.RIGHT)
for i,code in enumerate(ORDER):
    yy=2.46+i*.223
    if i%2==0:box(s,.65,yy-.025,11.99,.224,WHITE)
    f=T['families'][code]
    vals=[f"{code}  {NAMES[code]}",f"{M['family_counts'][code]:,}",f"{f['total']:,}",
          f"{100*f['first_try']/f['total']:.1f}%" if f['total'] else "—",
          f"{100*f['solved']/f['total']:.1f}%" if f['total'] else "—",f"{f['sft_eligible']:,}"]
    for j,(x,w,val) in enumerate(zip(xs,ws,vals)):
        txt(s,val,x,yy,w,.21,10.8,TEAL if j==5 else INK,j==5,
            align=PP_ALIGN.LEFT if j==0 else PP_ALIGN.RIGHT)

# 18 — Practical use appendix.
s = base("Load pinned releases, then enforce the scientific contract", "Appendix B / Getting started",
         "Keep task and teacher versions explicit in every training and evaluation manifest.",source="tasks",
         notes="These snippets identify released datasets only; they do not invoke models. Task archives must be unpacked on the host; only prompt.json system/user content may be sent to the model. The bundled no-tools Harbor adapter is CoordinateNoToolsAgent. For SFT, use messages with the pinned Snowball template and enable_thinking=True, mask system/user tokens, and check exact full conversation length. The teacher dataset uses v1.2.0 task prompts and rewards; reconcile this in a new release before claiming a v1.3 teacher run. Full immutable revisions appear in the following slide and in evidence.json.")
code='from datasets import load_dataset\n\ntasks = load_dataset(\n    "open-athena/pdbthink-coordinate-tasks",\n    split="train", revision="3734406cb97b1702844319f9a5d860cbbf8fe660"\n)\n\ntraces = load_dataset(\n    "open-athena/pdbthink-glm53-teacher-traces", "sft",\n    split="train", revision="91baf80abe4cd10e3258bbcc0addafd4d3d32b22"\n)'
box(s,.65,2.25,8.24,3.42,NIGHT,radius=True)
txt(s,code,.87,2.46,7.86,3.02,12,WHITE,font=MONO)
callout(s,"For evaluation / RL","Read prompt.json; run the bundled verifier outside the model.",9.12,2.25,3.52,1.82)
callout(s,"For SFT","Use messages, assistant-only loss and exact token-length checks.",9.12,4.29,3.52,1.81)
txt(s,"Trace configurations: sft = selected targets    ·    attempts = all scored responses    ·    outcomes = all cohort tasks",.75,6.43,11.9,.38,13,MUTED)

# 19 — Source ledger.
s = base("Pinned evidence and version boundaries", "Appendix C / Sources",
         "The deck is reproducible from evidence.json; full references are also in the speaker notes.",source="tasks",
         notes="Task release: "+URL['tasks']+"\nManifest: "+URL['manifest']+"\nContext audit: "+URL['context']+"\nCorrection evidence: "+URL['corrections']+"\nTeacher release: "+URL['teacher']+"\nTeacher summary: "+URL['teacher_summary']+"\nGenerator code: "+URL['code']+"\nSnowball tokenizer revision: cfc1d845dae89b067cdc7250d0164abefa5a69cf. The teacher's immutable served weight revision is unavailable. All figures refer to these pins rather than a moving main branch.")
sources=[("TASKS", "PDBThink Coordinate Tasks · v1.3.0", "3734406cb97b1702844319f9a5d860cbbf8fe660", "tasks", "100,000 tasks · prompt v5 · scorer 1.1.0 · definitions v1.0.0"),
         ("TEACHER", "GLM-5.3 Teacher Traces · v1.0.0", "91baf80abe4cd10e3258bbcc0addafd4d3d32b22", "teacher", "Historical prompts and rewards from coordinate tasks v1.2.0"),
         ("SNOWBALL", "Snowball-67B-A2B-5.7T-Mixed-RLVR-Step38", "cfc1d845dae89b067cdc7250d0164abefa5a69cf", "context", "Native tokenizer and thinking-enabled template · 32,768-token window")]
for i,(tag,title,sha,key,desc) in enumerate(sources):
    yy=2.19+i*1.26
    box(s,.65,yy,12.0,1.08,WHITE,radius=True)
    pill(s,tag,.85,yy+.18,1.17)
    txt(s,title,2.23,yy+.13,9.98,.34,18,TEAL,True,link=URL[key])
    txt(s,sha,2.24,yy+.55,9.97,.24,10.5,MUTED,font=MONO)
    txt(s,desc,2.24,yy+.83,9.97,.22,11,INK)
txt(s,"More evidence: manifest · native-context audit · correction audit · teacher summary · frozen generator code",.85,6.25,11.78,.42,14,MUTED)
txt(s,"Links are clickable in the deck and PDF.",.85,6.68,11.78,.25,11,TEAL)


prs.save(OUT)
(ROOT / "slide_notes.json").write_text(json.dumps(SLIDE_META,indent=2)+"\n")
print(f"Saved {OUT} ({len(prs.slides)} slides)")
